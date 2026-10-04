#!/usr/bin/env python
"""Prepare the Online Retail II dataset: download, convert to CSV, profile it and sample it.

Implements docs/10 section 2. Four subcommands, meant to be run in order:

    download  fetch both UCI archives into data/raw/ (git-ignored)
    convert   turn the .xlsx sheets into canonical CSV, also in data/raw/
    profile   measure the CSV and write docs/perfilado-dataset.md
    samples   build the baseline and natural/drift samples (steps 4-6)

The profiling report is the evidence used to decide the country-catalogue exclusions, the
alias list and the final field rules before sales_transaction.v1.yaml is frozen, in the
mandatory order of docs/10 section 2, step 3 bis. The samples in data/samples/ are built
afterwards, once the schema exists.

The baseline filter reads the frozen YAML but is not a validator: it only keeps rows it is
certain about. The authority on validation is the Java validator in pipeline-service (D40).
"""

import argparse
import csv
import random
import re
import shutil
import sys
import urllib.request
import zipfile
from collections import Counter
from collections.abc import Callable, Iterator
from contextlib import ExitStack
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

import yaml
from openpyxl import load_workbook

REPO_ROOT = Path(__file__).resolve().parents[1]
RAW_DIR = REPO_ROOT / "data" / "raw"
SAMPLES_DIR = REPO_ROOT / "data" / "samples"
SCHEMA_DIR = REPO_ROOT / "pipeline-service" / "src" / "main" / "resources" / "schemas"
SCHEMA_FILE = SCHEMA_DIR / "sales_transaction.v1.yaml"
PROFILE_REPORT = REPO_ROOT / "docs" / "perfilado-dataset.md"

HEADERS = [
    "Invoice",
    "StockCode",
    "Description",
    "Quantity",
    "InvoiceDate",
    "Price",
    "Customer ID",
    "Country",
]

LEGACY_HEADERS = [
    "InvoiceNo",
    "StockCode",
    "Description",
    "Quantity",
    "InvoiceDate",
    "UnitPrice",
    "CustomerID",
    "Country",
]


class Source:
    def __init__(self, key: str, url: str, xlsx: str, csv_name: str, headers: list[str]) -> None:
        self.key = key
        self.url = url
        self.xlsx = RAW_DIR / xlsx
        self.csv = RAW_DIR / csv_name
        self.headers = headers


CURRENT = Source(
    key="online-retail-ii",
    url="https://archive.ics.uci.edu/static/public/502/online+retail+ii.zip",
    xlsx="online_retail_ii.xlsx",
    csv_name="online_retail_ii.csv",
    headers=HEADERS,
)

LEGACY = Source(
    key="online-retail-legacy",
    url="https://archive.ics.uci.edu/static/public/352/online+retail.zip",
    xlsx="online_retail_legacy.xlsx",
    csv_name="online_retail_legacy.csv",
    headers=LEGACY_HEADERS,
)
SOURCES = [CURRENT, LEGACY]


def download(force: bool = False) -> None:
    RAW_DIR.mkdir(parents=True, exist_ok=True)
    for source in SOURCES:
        if source.xlsx.exists() and not force:
            print(f"  {source.xlsx.name}: already present, skipping")
            continue
        archive = RAW_DIR / f"{source.key}.zip"
        print(f"  downloading {source.url}")
        try:
            # S310 wants the scheme validated; source.url is a hardcoded https constant above.
            urllib.request.urlretrieve(source.url, archive)  # noqa: S310
        except OSError as exc:
            print(
                f"\nERROR: could not download {source.url}\n  {exc}\n\n"
                f"Download it by hand, unzip it and place the .xlsx at {source.xlsx}, "
                f"then re-run `convert`.",
                file=sys.stderr,
            )
            raise SystemExit(1) from exc

        with zipfile.ZipFile(archive) as zf:
            members = [n for n in zf.namelist() if n.lower().endswith(".xlsx")]
            if len(members) != 1:
                raise SystemExit(f"expected exactly one .xlsx in {archive}, found {members}")
            with zf.open(members[0]) as src, source.xlsx.open("wb") as dst:
                shutil.copyfileobj(src, dst)
        archive.unlink()
        print(f"  {source.xlsx.name}: {source.xlsx.stat().st_size} bytes")


def _cell_to_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        dec = Decimal(repr(value))
        if dec == dec.to_integral_value():
            return str(dec.to_integral_value())
        return format(dec.normalize(), "f")
    return str(value)


def _convert_source(source: Source) -> int:
    if not source.xlsx.exists():
        raise SystemExit(f"{source.xlsx} not found - run `download` first")

    workbook = load_workbook(source.xlsx, read_only=True, data_only=True)
    written = 0
    # newline="" is required by the csv module; \n keeps the line endings stable across platforms
    # and utf-8 (not utf-8-sig) keeps the file free of a BOM, which would fake a schema drift
    # on the first header (AC-01.7).
    with source.csv.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, lineterminator="\n", quoting=csv.QUOTE_MINIMAL)
        writer.writerow(source.headers)
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            rows = sheet.iter_rows(values_only=True)
            header_row = next(rows, None)
            if header_row is None:
                continue
            found = [str(cell).strip() if cell is not None else "" for cell in header_row]
            if found != source.headers:
                raise SystemExit(
                    f"unexpected headers in sheet {sheet_name!r} of {source.xlsx.name}:\n"
                    f"  found:    {found}\n  expected: {source.headers}\n"
                    "The schema is pinned to these headers (docs/02 section 2); a change here is "
                    "a dataset change and needs an owner decision, not a silent remap."
                )
            for row in rows:
                if all(cell is None for cell in row):
                    continue
                writer.writerow([_cell_to_text(cell) for cell in row])
                written += 1
        workbook.close()
    print(f"  {source.csv.name}: {written} data rows")
    return written


def convert() -> None:
    for source in SOURCES:
        _convert_source(source)


INVOICE_RE = re.compile(r"^[CA]?\d{6}$")
STOCK_CODE_RE = re.compile(r"^[A-Za-z0-9](?:[A-Za-z0-9_ \-]{0,18}[A-Za-z0-9_])?$")
CUSTOMER_ID_RE = re.compile(r"^\d{5}$")
NULL_TOKENS = {"", "N/A", "NA", "null", "NULL", "None", "-", "?"}
DATE_FORMAT = "%Y-%m-%d %H:%M:%S"
MIN_DATE = datetime(2009, 1, 1)


class Profile:
    def __init__(self) -> None:
        self.rows = 0
        self.empty: Counter[str] = Counter()
        self.null_token: Counter[str] = Counter()
        self.countries: Counter[str] = Counter()
        self.invoice_bad: Counter[str] = Counter()
        self.invoice_prefix: Counter[str] = Counter()
        self.stock_bad: Counter[str] = Counter()
        self.stock_lengths: Counter[int] = Counter()
        self.stock_seen: dict[str, set[str]] = {}
        self.customer_bad: Counter[str] = Counter()
        self.price_scale: Counter[int] = Counter()
        self.price_negative = 0
        self.quantity_zero = 0
        self.quantity_min: int | None = None
        self.quantity_max: int | None = None
        self.price_min: Decimal | None = None
        self.price_max: Decimal | None = None
        self.date_min: datetime | None = None
        self.date_max: datetime | None = None
        self.description_over_255 = 0
        self.description_untrimmed = 0
        self.br01_violations = 0
        self.br02_violations = 0
        self.unparsable: Counter[str] = Counter()

    def feed(self, row: dict[str, str]) -> None:
        self.rows += 1

        for field, value in row.items():
            if value == "":
                self.empty[field] += 1
            elif value.strip() in NULL_TOKENS:
                self.null_token[field] += 1

        invoice = row["Invoice"]
        if not INVOICE_RE.match(invoice):
            self.invoice_bad[invoice] += 1
        self.invoice_prefix[invoice[:1] if invoice[:1].isalpha() else "(digit)"] += 1

        stock = row["StockCode"]
        self.stock_lengths[len(stock)] += 1
        if not STOCK_CODE_RE.match(stock):
            self.stock_bad[stock] += 1
        # Codes that differ only in case are different products in this dataset, which is why
        # NORMALIZE_CASE is forbidden on stock_code (ADR-0008).
        self.stock_seen.setdefault(stock.casefold(), set()).add(stock)

        description = row["Description"]
        if len(description) > 255:
            self.description_over_255 += 1
        if description != description.strip():
            self.description_untrimmed += 1

        quantity: int | None = None
        try:
            quantity = int(row["Quantity"])
        except ValueError:
            self.unparsable["Quantity"] += 1
        else:
            if quantity == 0:
                self.quantity_zero += 1
            self.quantity_min = (
                quantity if self.quantity_min is None else min(self.quantity_min, quantity)
            )
            self.quantity_max = (
                quantity if self.quantity_max is None else max(self.quantity_max, quantity)
            )

        try:
            price = Decimal(row["Price"])
        except ArithmeticError:
            self.unparsable["Price"] += 1
        else:
            # Decimal("NaN") and Decimal("Infinity") parse without raising, and their exponent is
            # a string ("n", "F") instead of an int, so negating it would abort the whole run with
            # a TypeError that `except ArithmeticError` never sees. Comparisons against NaN are
            # also meaningless, which would silently corrupt the min/max range. A non-finite price
            # is not a value the schema can accept, so it counts as unparsable.
            exponent = price.as_tuple().exponent
            if not isinstance(exponent, int):
                self.unparsable["Price"] += 1
            else:
                self.price_scale[max(0, -exponent)] += 1
                if price < 0:
                    self.price_negative += 1
                self.price_min = price if self.price_min is None else min(self.price_min, price)
                self.price_max = price if self.price_max is None else max(self.price_max, price)

        try:
            moment = datetime.strptime(row["InvoiceDate"], DATE_FORMAT)
        except ValueError:
            self.unparsable["InvoiceDate"] += 1
        else:
            self.date_min = moment if self.date_min is None else min(self.date_min, moment)
            self.date_max = moment if self.date_max is None else max(self.date_max, moment)

        customer = row["Customer ID"]
        if customer != "" and not CUSTOMER_ID_RE.match(customer):
            self.customer_bad[customer] += 1

        self.countries[row["Country"]] += 1

        # BR-01 / BR-02 are only meaningful when both fields parsed.
        if quantity is not None and INVOICE_RE.match(invoice):
            if invoice.startswith("C") and quantity >= 0:
                self.br01_violations += 1
            elif not invoice.startswith("C") and quantity <= 0:
                self.br02_violations += 1


def _pct(part: int, total: int) -> str:
    return f"{(100 * part / total):.2f} %" if total else "n/a"


def profile() -> None:
    if not CURRENT.csv.exists():
        raise SystemExit(f"{CURRENT.csv} not found - run `convert` first")

    data = Profile()
    with CURRENT.csv.open(encoding="utf-8", newline="") as handle:
        for row in csv.DictReader(handle):
            data.feed(row)

    case_collisions = {k: v for k, v in data.stock_seen.items() if len(v) > 1}
    lines: list[str] = []
    add = lines.append

    add("# Perfilado del dataset — Online Retail II")
    add("")
    add(
        f"Generado por `tools/download_dataset.py profile` el "
        f"{datetime.now().strftime('%Y-%m-%d')} sobre `data/raw/{CURRENT.csv.name}` "
        f"({data.rows} filas de datos)."
    )
    add("")
    add(
        "> Este documento es la **evidencia** con la que se deciden las exclusiones del catálogo "
        "de países, la lista de alias y las reglas finales de campo antes de congelar "
        "`sales_transaction.v1.yaml`, en el orden obligatorio de `10` §2 paso 3 bis. "
        "Los números no se editan a mano: se regeneran con el script (`09` §2.3)."
    )
    add("")

    add("## 1. Cabeceras reales")
    add("")
    add(f"Hojas: `Year 2009-2010` y `Year 2010-2011`. Cabeceras: `{'`, `'.join(HEADERS)}`.")
    add("")
    add(
        "Coinciden exactamente con `02` §2. La ficha de UCI publica los nombres de la versión "
        f"anterior (`{'`, `'.join(LEGACY_HEADERS)}`), que es justamente el caso de *schema "
        "drift* de FR-05; esa versión se conserva en "
        f"`data/raw/{LEGACY.csv.name}` para construir `drift-legacy-headers.csv`."
    )
    add("")

    add("## 2. Ausencias por columna")
    add("")
    add("| Columna | Vacías | % | Token nulo (`N/A`, `-`, ...) |")
    add("|---|---:|---:|---:|")
    for field in HEADERS:
        add(
            f"| `{field}` | {data.empty[field]} | {_pct(data.empty[field], data.rows)} "
            f"| {data.null_token[field]} |"
        )
    add("")

    add("## 3. `Invoice`")
    add("")
    add(
        f"- Filas que cumplen `^[CA]?\\d{{6}}$`: {data.rows - sum(data.invoice_bad.values())} "
        f"({_pct(data.rows - sum(data.invoice_bad.values()), data.rows)})."
    )
    add(
        f"- Filas que no lo cumplen: {sum(data.invoice_bad.values())} "
        f"({len(data.invoice_bad)} valores distintos)."
    )
    if data.invoice_bad:
        sample = ", ".join(f"`{v}` x{n}" for v, n in data.invoice_bad.most_common(10))
        add(f"- Ejemplos: {sample}.")
    add(
        "- Prefijos observados: "
        + ", ".join(f"`{p}` x{n}" for p, n in sorted(data.invoice_prefix.items()))
        + "."
    )
    add("")

    add("## 4. `StockCode`")
    add("")
    add(f"- Longitudes: mín. {min(data.stock_lengths)}, máx. {max(data.stock_lengths)}.")
    add(
        f"- Filas que no cumplen la regex propuesta: {sum(data.stock_bad.values())} "
        f"({len(data.stock_bad)} valores distintos)."
    )
    if data.stock_bad:
        sample = ", ".join(f"`{v}` x{n}" for v, n in data.stock_bad.most_common(15))
        add(f"- Ejemplos: {sample}.")
    add("")
    add(
        f"**Códigos que solo difieren en mayúsculas/minúsculas: {len(case_collisions)} grupos.** "
        "Es la evidencia de que `NORMALIZE_CASE` está prohibido en `stock_code` (ADR-0008): "
        "cambiar la caja cambia de producto."
    )
    if case_collisions:
        sample = "; ".join(
            " / ".join(sorted(f"`{s}`" for s in variants))
            for variants in list(case_collisions.values())[:10]
        )
        add(f"Ejemplos: {sample}.")
    add("")

    add("## 5. `Quantity` y `Price`")
    add("")
    add(
        f"- `Quantity`: rango [{data.quantity_min}, {data.quantity_max}]; "
        f"valores 0: {data.quantity_zero}."
    )
    add(f"- `Price`: rango [{data.price_min}, {data.price_max}]; negativos: {data.price_negative}.")
    add(
        "- Decimales de `Price`: "
        + ", ".join(f"{k} → {v} filas" for k, v in sorted(data.price_scale.items()))
        + "."
    )
    over_two = sum(v for k, v in data.price_scale.items() if k > 2)
    add(
        f"- Filas con más de 2 decimales en `Price`: {over_two}. "
        "Con `decimal(12,2)` quedan fuera del baseline: el schema exige máximo 2 decimales y "
        "`02` §2 prohíbe redondear (sería `INVALID_FORMAT`)."
    )
    add("")

    add("## 6. `InvoiceDate`")
    add("")
    add(f"- Rango: {data.date_min} … {data.date_max}.")
    add(
        f"- Filas anteriores al mínimo del schema ({MIN_DATE:%Y-%m-%d}): "
        f"{'0' if data.date_min and data.date_min >= MIN_DATE else 'revisar'}."
    )
    add("")

    add("## 7. `Customer ID`")
    add("")
    add(f"- Vacías: {data.empty['Customer ID']} ({_pct(data.empty['Customer ID'], data.rows)}).")
    add(f"- No vacías que no cumplen `^\\d{{5}}$`: {sum(data.customer_bad.values())}.")
    if data.customer_bad:
        sample = ", ".join(f"`{v}` x{n}" for v, n in data.customer_bad.most_common(10))
        add(f"- Ejemplos: {sample}.")
    add("")

    add("## 8. `Description`")
    add("")
    add(f"- Más de 255 caracteres: {data.description_over_255}.")
    add(
        f"- Con espacios al inicio o final: {data.description_untrimmed} "
        "(suciedad natural; el MVP no normaliza texto libre, `02` §2)."
    )
    add("")

    add("## 9. Reglas de negocio")
    add("")
    add(f"- `BR-01` (factura `C...` debe tener `quantity < 0`) incumplida: {data.br01_violations}.")
    add(
        f"- `BR-02` (factura sin `C` debe tener `quantity > 0`) incumplida: {data.br02_violations}."
    )
    add("")

    add("## 10. Catálogo de países")
    add("")
    add(
        f"Valores distintos de `Country` en el dataset completo (ambas hojas): "
        f"**{len(data.countries)}**."
    )
    add("")
    add("| País | Filas | % |")
    add("|---|---:|---:|")
    for country, count in sorted(data.countries.items(), key=lambda kv: (-kv[1], kv[0])):
        add(f"| `{country}` | {count} | {_pct(count, data.rows)} |")
    add("")
    add("### Exclusiones propuestas")
    add("")
    add(
        "`02` §2 define `countries.v1.txt` como estos valores distintos **menos** las exclusiones "
        "decididas aquí, cada una con su motivo. Candidatos nombrados por el documento:"
    )
    add("")
    add("| Candidato | Filas | Motivo propuesto |")
    add("|---|---:|---|")
    for candidate in ("Unspecified", "European Community"):
        candidate_rows = data.countries.get(candidate)
        present = str(candidate_rows) if candidate_rows is not None else "no aparece"
        add(
            f"| `{candidate}` | {present} | No es un país: es un marcador de dato ausente o un "
            "bloque supranacional. Si entrara al catálogo, una fila sin país real pasaría la "
            "validación como válida en lugar de ir a cuarentena. |"
        )
    add("")

    if data.unparsable:
        add("## 11. Celdas no interpretables")
        add("")
        for field, count in data.unparsable.items():
            add(f"- `{field}`: {count}.")
        add("")

    PROFILE_REPORT.parent.mkdir(parents=True, exist_ok=True)
    PROFILE_REPORT.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(f"  {PROFILE_REPORT.relative_to(REPO_ROOT)}: {len(lines)} lines")
    print(
        f"  rows={data.rows} countries={len(data.countries)} "
        f"stock_case_collisions={len(case_collisions)}"
    )


SEED = 42
CLEAN_TARGETS = {
    1_000: SAMPLES_DIR / "clean-1k.csv",
    10_000: RAW_DIR / "clean-10k.csv",
    100_000: RAW_DIR / "clean-100k.csv",
}
RAW_NATURAL_TARGET = SAMPLES_DIR / "raw-natural-2k.csv"
RAW_NATURAL_SIZE = 2_000
DRIFT_TARGET = SAMPLES_DIR / "drift-legacy-headers.csv"
DRIFT_SIZE = 1_000

SCHEMA_KEYS = {"name", "version", "nullTokens", "fields", "businessRules", "derived"}
FIELD_KEYS = {
    "name",
    "sourceHeader",
    "type",
    "required",
    "pattern",
    "maxLength",
    "min",
    "max",
    "notEqualTo",
    "format",
    "scale",
    "decimalSeparator",
    "catalogue",
    "caseSensitive",
}
BUSINESS_RULES = {"BR-01", "BR-02"}
# Canonical shapes only. re.ASCII: Python's \d also matches non-ASCII digits, Java's does not.
INTEGER_SHAPE = re.compile(r"-?[1-9]\d*", re.ASCII)
DATE_FORMATS = {
    "yyyy-MM-dd HH:mm:ss": (
        re.compile(r"\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2}", re.ASCII),
        "%Y-%m-%d %H:%M:%S",
    ),
}


def _value_check(spec: dict[str, Any]) -> Callable[[str], bool]:
    kind = spec["type"]
    if kind == "string":
        pattern = re.compile(spec["pattern"], re.ASCII) if "pattern" in spec else None
        max_length = int(spec["maxLength"]) if "maxLength" in spec else None

        def check_string(value: str) -> bool:
            # fullmatch, not match: Python's $ also matches before a trailing newline.
            if pattern is not None and pattern.fullmatch(value) is None:
                return False
            # Java's String.length() counts UTF-16 units, never fewer than Python's code points.
            return max_length is None or len(value.encode("utf-16-le")) // 2 <= max_length

        return check_string

    if kind == "integer":
        low, high = int(spec["min"]), int(spec["max"])
        excluded = int(spec["notEqualTo"]) if "notEqualTo" in spec else None

        def check_integer(value: str) -> bool:
            if INTEGER_SHAPE.fullmatch(value) is None:
                return False
            number = int(value)
            return low <= number <= high and number != excluded

        return check_integer

    if kind == "datetime":
        if spec["format"] not in DATE_FORMATS:
            raise SystemExit(f"{spec['name']}: unsupported date format {spec['format']!r}")
        shape, layout = DATE_FORMATS[spec["format"]]
        earliest = datetime.strptime(str(spec["min"]), layout)
        latest = datetime.now() if spec["max"] == "now" else datetime.strptime(spec["max"], layout)

        def check_datetime(value: str) -> bool:
            # strptime alone accepts "2010-1-5 8:05:00"; the shape check keeps it canonical.
            if shape.fullmatch(value) is None:
                return False
            try:
                moment = datetime.strptime(value, layout)
            except ValueError:
                return False
            return earliest <= moment <= latest

        return check_datetime

    if kind == "decimal":
        if spec["decimalSeparator"] != ".":
            raise SystemExit(f"{spec['name']}: unsupported decimal separator")
        shape = re.compile(rf"-?(?:0|[1-9]\d*)(?:\.\d{{1,{int(spec['scale'])}}})?", re.ASCII)
        low_price, high_price = Decimal(str(spec["min"])), Decimal(str(spec["max"]))

        def check_decimal(value: str) -> bool:
            return shape.fullmatch(value) is not None and low_price <= Decimal(value) <= high_price

        return check_decimal

    if kind == "enum":
        lines = (SCHEMA_DIR / spec["catalogue"]).read_text(encoding="utf-8").splitlines()
        catalogue = frozenset(line for line in lines if line)
        return lambda value: value in catalogue

    raise SystemExit(f"{spec['name']}: unsupported type {kind!r}")


class BaselineFilter:
    """Keeps a row only if it is certain the row passes schema v1; any doubt drops it (D40)."""

    def __init__(self, schema_file: Path) -> None:
        schema = yaml.safe_load(schema_file.read_text(encoding="utf-8"))
        # Fail on anything this filter does not implement, instead of silently ignoring a rule.
        if set(schema) - SCHEMA_KEYS:
            raise SystemExit(f"unknown schema keys: {sorted(set(schema) - SCHEMA_KEYS)}")
        if {rule["code"] for rule in schema["businessRules"]} != BUSINESS_RULES:
            raise SystemExit("business rules differ from the BR-01/BR-02 this filter implements")
        for spec in schema["fields"]:
            if set(spec) - FIELD_KEYS or spec.get("caseSensitive", True) is not True:
                raise SystemExit(f"{spec['name']}: rule not supported by the baseline filter")

        self.null_tokens = frozenset(schema["nullTokens"])
        self.headers: list[str] = [spec["sourceHeader"] for spec in schema["fields"]]
        self.checks = [(bool(spec["required"]), _value_check(spec)) for spec in schema["fields"]]
        names = [spec["name"] for spec in schema["fields"]]
        self.invoice = names.index("invoice_no")
        self.quantity = names.index("quantity")

    def accepts(self, row: list[str]) -> bool:
        if len(row) != len(self.checks):
            return False
        for (required, check), value in zip(self.checks, row, strict=True):
            if value.strip(" \t") in self.null_tokens:
                if required:
                    return False
                if value == "":
                    continue
            if not check(value):
                return False
        # BR-01 / BR-02 only carry a description in the YAML, so they are coded here as in Java.
        quantity = int(row[self.quantity])
        return quantity < 0 if row[self.invoice].startswith("C") else quantity > 0


def _seeded() -> random.Random:
    # S311 is about cryptography; here the fixed seed is the point (AC-19.3).
    return random.Random(SEED)  # noqa: S311


def _read_rows(path: Path) -> Iterator[list[str]]:
    with path.open(encoding="utf-8", newline="") as handle:
        yield from csv.reader(handle)


def _write_samples(source: Path, picks: dict[Path, set[int]]) -> None:
    """Copy the chosen data rows (0-based) of source into each target, in file order."""
    rows = _read_rows(source)
    header = next(rows)
    with ExitStack() as stack:
        writers = []
        for target, chosen in picks.items():
            target.parent.mkdir(parents=True, exist_ok=True)
            handle = stack.enter_context(target.open("w", encoding="utf-8", newline=""))
            writer = csv.writer(handle, lineterminator="\n", quoting=csv.QUOTE_MINIMAL)
            writer.writerow(header)
            writers.append((chosen, writer))
        for index, row in enumerate(rows):
            for chosen, writer in writers:
                if index in chosen:
                    writer.writerow(row)
    for target, chosen in picks.items():
        print(f"  {target.relative_to(REPO_ROOT)}: {len(chosen)} rows")


def samples() -> None:
    for source in SOURCES:
        if not source.csv.exists():
            raise SystemExit(f"{source.csv} not found - run `convert` first")

    baseline = BaselineFilter(SCHEMA_FILE)
    rows = _read_rows(CURRENT.csv)
    if next(rows) != baseline.headers:
        raise SystemExit(f"{CURRENT.csv.name} headers do not match {SCHEMA_FILE.name}")
    total = 0
    clean: list[int] = []
    for index, row in enumerate(rows):
        total += 1
        if baseline.accepts(row):
            clean.append(index)
    print(f"  baseline: {len(clean)} of {total} rows pass schema v1")

    # One draw sliced three ways, so clean-1k is a subset of clean-10k and of clean-100k.
    drawn = _seeded().sample(clean, max(CLEAN_TARGETS))
    picks = {target: set(drawn[:size]) for size, target in CLEAN_TARGETS.items()}
    picks[RAW_NATURAL_TARGET] = set(_seeded().sample(range(total), RAW_NATURAL_SIZE))
    _write_samples(CURRENT.csv, picks)

    legacy_total = sum(1 for _ in _read_rows(LEGACY.csv)) - 1
    drift = set(_seeded().sample(range(legacy_total), DRIFT_SIZE))
    _write_samples(LEGACY.csv, {DRIFT_TARGET: drift})


# --------------------------------------------------------------------------------------


def main() -> None:
    # A literal, not __doc__: python -OO strips docstrings, which would make __doc__ None here.
    parser = argparse.ArgumentParser(
        description="Prepare the Online Retail II dataset (docs/10 section 2)."
    )
    sub = parser.add_subparsers(dest="command", required=True)
    download_parser = sub.add_parser("download", help="fetch both UCI archives into data/raw/")
    download_parser.add_argument(
        "--force", action="store_true", help="re-download even if the .xlsx is already there"
    )
    sub.add_parser("convert", help="convert the .xlsx sheets to canonical CSV")
    sub.add_parser("profile", help="write docs/perfilado-dataset.md from the CSV")
    sub.add_parser("samples", help="build the baseline, natural and drift samples (seed 42)")
    sub.add_parser("all", help="download, convert, profile and sample in order")

    args = parser.parse_args()
    if args.command == "download":
        download(force=args.force)
    elif args.command == "convert":
        convert()
    elif args.command == "profile":
        profile()
    elif args.command == "samples":
        samples()
    else:
        download()
        convert()
        profile()
        samples()


if __name__ == "__main__":
    main()
